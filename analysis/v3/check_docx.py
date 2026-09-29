# -*- coding: utf-8 -*-
"""Sanity check of a generated .docx: counts, leftover LaTeX, and that key numbers survived."""
import sys, re
import docx

d = docx.Document(sys.argv[1])
paras = "\n".join(p.text for p in d.paragraphs)
cells = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
allt = paras + "\n" + cells
print("paragraphs", len(d.paragraphs), "tables", len(d.tables), "images", len(d.inline_shapes))
bad = re.findall(r"\\[a-zA-Z]+|\?\?|\{|\}", allt)
print("leftover LaTeX tokens:", len(bad), sorted(set(bad))[:20])
for key in ("+20.72", "-58.32", "-3.20", "+15.31", "-24.99", "-100.18", "1,500", "249"):
    print("%-8s %s" % (key, "found" if key in allt else "MISSING"))
print("abstract starts:", paras[paras.find("The guidance hypothesis holds"):][:80])
